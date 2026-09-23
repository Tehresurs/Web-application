from datetime import datetime
from html import escape
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.services.employee_matcher import find_employee_by_full_name
from app.services.object_matcher import match_object
from app.reports.intermediate_report import build_intermediate_report

INPUT_FOLDER = Path("input")
OUTPUT_FOLDER = Path("output")
HTML_REPORT = OUTPUT_FOLDER / "parser_diagnostic.html"


def clean_text(text: str) -> str:
    if not text:
        return ""
    return "\n".join(
        cleaned
        for line in text.splitlines()
        if (cleaned := " ".join(line.split()))
    )


def normalize_for_search(text: str) -> str:
    return clean_text(text).lower().replace("ё", "е")


def get_xml_text(element) -> str:
    return clean_text("".join(
        text_element.text or ""
        for text_element in element.iter(qn("w:t"))
    ))


def get_table_rows(table):
    result = []
    for row_index, tr in enumerate(table._tbl.findall(qn("w:tr"))):
        cells = []
        for child in tr:
            if child.tag == qn("w:tc"):
                cells.append({"text": get_xml_text(child), "source_type": "tc"})
            elif child.tag == qn("w:sdt"):
                tc_elements = list(child.iter(qn("w:tc")))
                if tc_elements:
                    cells.extend(
                        {"text": get_xml_text(tc), "source_type": "sdt"}
                        for tc in tc_elements
                    )
                else:
                    text = get_xml_text(child)
                    if text:
                        cells.append({"text": text, "source_type": "sdt"})
        for cell_index, cell in enumerate(cells):
            cell["row_index"] = row_index
            cell["cell_index"] = cell_index
        result.append({"row_index": row_index, "cells": cells})
    return result


def make_result(raw=None, value=None, found=False, table_index=None,
                row_index=None, cell_index=None, source_type=None):
    return {
        "found": found,
        "raw": raw,
        "value": value,
        "source": {
            "table": table_index,
            "row": row_index,
            "cell": cell_index,
            "type": source_type,
        } if found else None,
    }


def next_non_empty_cell(row, start_index):
    for cell in row["cells"][start_index + 1:]:
        if clean_text(cell["text"]):
            return cell
    return None


def find_value_after_label(tables, labels):
    normalized_labels = {normalize_for_search(label) for label in labels}
    for table in tables:
        for row in table["rows"]:
            for index, cell in enumerate(row["cells"]):
                if normalize_for_search(cell["text"]) in normalized_labels:
                    value_cell = next_non_empty_cell(row, index)
                    if value_cell:
                        raw = clean_text(value_cell["text"])
                        return make_result(
                            raw, raw, True, table["table_index"],
                            value_cell["row_index"], value_cell["cell_index"],
                            value_cell["source_type"],
                        )
    return make_result()


def find_object(tables):
    return find_value_after_label(tables, ["Объект"])


def find_general_contractor(tables):
    return find_value_after_label(tables, ["Генподрядчик"])


def normalize_date(raw):
    if not raw:
        return None
    for date_format in ("%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(raw.strip(), date_format).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def find_report_date(tables):
    for table in tables:
        for row in table["rows"]:
            cells = row["cells"]
            normalized = [normalize_for_search(cell["text"]) for cell in cells]
            if "дата" not in normalized or "направление контроля" not in normalized:
                continue
            for index, cell in enumerate(cells):
                if normalize_for_search(cell["text"]) == "дата":
                    value_cell = next_non_empty_cell(row, index)
                    if value_cell:
                        raw = clean_text(value_cell["text"])
                        value = normalize_date(raw)
                        if value:
                            return make_result(
                                raw, value, True, table["table_index"],
                                value_cell["row_index"], value_cell["cell_index"],
                                value_cell["source_type"],
                            )
    return make_result()


def find_subcontractor(tables):
    for table in tables:
        for row in table["rows"]:
            for index, cell in enumerate(row["cells"]):
                if normalize_for_search(cell["text"]) != "субподрядчик":
                    continue
                for value_cell in row["cells"][index + 1:]:
                    raw = clean_text(value_cell["text"])
                    if raw:
                        return make_result(
                            raw, raw, True, table["table_index"],
                            value_cell["row_index"], value_cell["cell_index"],
                            value_cell["source_type"],
                        )
                return make_result()
    return make_result()


def looks_like_full_name(text):
    words = text.split() if text else []
    return len(words) == 3 and all(
        word.replace("-", "").replace("ё", "е").replace("Ё", "Е").isalpha()
        for word in words
    )


def is_project_manager(text: str) -> bool:
    normalized = normalize_for_search(text)
    return "руководитель проекта" in normalized or normalized == "руководитель"


def find_full_name_in_rows(rows, start_position, max_rows=3):
    end_position = min(start_position + 1 + max_rows, len(rows))
    for row_position in range(start_position + 1, end_position):
        for cell in rows[row_position]["cells"]:
            text = clean_text(cell["text"])
            if looks_like_full_name(text):
                return cell
    return None


def find_inspector(tables):
    if not tables:
        return {
            "found": False,
            "status": "not_found",
            "full_name": None,
            "source": None,
        }

    table = tables[0]
    rows = table["rows"]

    if len(rows) < 2:
        return {
            "found": False,
            "status": "not_found",
            "full_name": None,
            "source": None,
        }

    inspector_row = rows[-2]

    if not inspector_row["cells"]:
        return {
            "found": False,
            "status": "not_found",
            "full_name": None,
            "source": None,
        }

    name_cell = inspector_row["cells"][0]
    full_name = clean_text(name_cell["text"])

    source = {
        "table": table["table_index"],
        "row": inspector_row["row_index"],
        "cell": name_cell["cell_index"],
        "type": name_cell["source_type"],
    }

    if not looks_like_full_name(full_name):
        return {
            "found": False,
            "status": "invalid_name",
            "full_name": {
                "raw": full_name,
                "value": full_name,
            },
            "source": source,
        }

    return {
        "found": True,
        "status": "found",
        "full_name": {
            "raw": full_name,
            "value": full_name,
        },
        "source": source,
    }


def extract_business_data(tables):
    inspector = find_inspector(tables)

    employee_match = {
        "status": "not_checked",
        "employee": None,
        "candidates": [],
    }

    if inspector["found"]:
        full_name = inspector["full_name"]["value"]
        employee_match = find_employee_by_full_name(full_name)

    object_data = find_object(tables)
    object_match = {
        "status": "not_checked",
        "object": None,
        "candidates": [],
    }

    if object_data["found"]:
        object_name = object_data["raw"] or object_data["value"]
        object_match = match_object(object_name)

    return {
        "report_date": find_report_date(tables),
        "object": object_data,
        "object_match": object_match,
        "inspector": inspector,
        "employee_match": employee_match,
        "general_contractor": find_general_contractor(tables),
        "subcontractor": find_subcontractor(tables),
    }


def read_document(file_path: Path):
    document = Document(file_path)
    paragraphs = [
        {"index": index, "text": text}
        for index, paragraph in enumerate(document.paragraphs)
        if (text := clean_text(paragraph.text))
    ]
    tables = [
        {"table_index": index, "rows": get_table_rows(table)}
        for index, table in enumerate(document.tables)
    ]
    return {
        "filename": file_path.name,
        "paragraphs": paragraphs,
        "tables": tables,
        "business_data": extract_business_data(tables),
    }


def status_html(field):
    return '<span class="ok">✓ найдено</span>' if field["found"] else '<span class="error">✗ не найдено</span>'


def value_html(field, display_raw=False):
    if not field["found"]:
        return '<span class="empty">—</span>'
    value = field["raw"] if display_raw else field["value"]
    return escape(str(value)).replace("\n", "<br>")


def source_html(field):
    if not field["found"]:
        return "—"
    source = field["source"]
    return f'T{source["table"]} / R{source["row"]} / C{source["cell"]} / {source["type"]}'


def inspector_status_html(inspector):
    status = inspector["status"]
    if status == "found":
        return '<span class="ok">✓ найден</span>'
    if status == "ambiguous":
        return '<span class="warning">⚠ несколько кандидатов</span>'
    return '<span class="error">✗ не найден</span>'


def build_summary_table(document):
    data = document["business_data"]
    rows = [
        ("Дата отчёта", data["report_date"]),
        ("Объект", data["object"]),
        ("Генподрядчик", data["general_contractor"]),
        ("Субподрядчик", data["subcontractor"]),
    ]
    html = '<h2>Распознанные данные</h2><table class="summary"><tr><th>Поле</th><th>Исходное значение</th><th>Значение парсера</th><th>Статус</th><th>Источник</th></tr>'
    for name, field in rows:
        html += f'<tr><td><strong>{escape(name)}</strong></td><td>{value_html(field, True)}</td><td>{value_html(field)}</td><td>{status_html(field)}</td><td>{source_html(field)}</td></tr>'
    return html + "</table>"


def build_employee_match_html(document):
    data = document["business_data"]
    inspector = data["inspector"]
    match = data["employee_match"]
    html = '<h3>Проверка сотрудника по базе</h3><table class="summary"><tr><th>Параметр</th><th>Значение</th></tr>'

    if not inspector["found"]:
        html += '<tr><td>Статус</td><td><span class="error">✗ ФИО НЕ РАСПОЗНАНО</span></td></tr>'
        return html + "</table>"

    word_name = inspector["full_name"]["value"]
    html += f'<tr><td>ФИО из Word</td><td>{escape(word_name)}</td></tr>'

    if match["status"] == "found":
        employee = match["employee"]
        position = employee["position_name"] or "—"
        phone = employee["phone"] or "—"
        html += '<tr><td>Статус</td><td><span class="ok">✓ СОТРУДНИК НАЙДЕН</span></td></tr>'
        html += (
            f'<tr><td>ID</td><td>{employee["id"]}</td></tr>'
            f'<tr><td>ФИО в базе</td><td>{escape(employee["full_name"])}</td></tr>'
            f'<tr><td>Должность из БД</td><td><strong>{escape(position)}</strong></td></tr>'
            f'<tr><td>Телефон</td><td>{escape(phone)}</td></tr>'
        )
    else:
        html += '<tr><td>Статус</td><td><span class="warning">⚠ СОТРУДНИК НЕ НАЙДЕН В БАЗЕ</span></td></tr><tr><td>Должность</td><td>—</td></tr>'

    return html + "</table>"


def build_global_summary(documents):
    html = '<h2>Сводка по всем отчётам</h2><table class="summary"><tr><th>Файл</th><th>Должность</th><th>ФИО специалиста</th><th>Статус</th></tr>'
    for document in documents:
        inspector = document["business_data"]["inspector"]
        filename = escape(document["filename"])
        employee_match = document["business_data"]["employee_match"]
        if inspector["status"] == "found":
            full_name = escape(inspector["full_name"]["value"])
        else:
            full_name = "—"
        if employee_match["status"] == "found":
            position = escape(employee_match["employee"]["position_name"] or "—")
        else:
            position = "—"
        html += f'<tr><td>{filename}</td><td>{position}</td><td>{full_name}</td><td>{inspector_status_html(inspector)}</td></tr>'
    return html + "</table>"


def build_html_report(documents):
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    html = '''<!DOCTYPE html><html lang="ru"><head><meta charset="UTF-8"><title>Проверка Word-парсера</title><style>
body{font-family:Arial,sans-serif;background:#f5f6f8;margin:0;padding:30px}h1{margin-bottom:30px}.document{background:white;border-radius:10px;padding:25px;margin-bottom:35px;box-shadow:0 2px 8px rgba(0,0,0,.08)}.filename{font-size:20px;font-weight:bold;margin-bottom:20px}table{border-collapse:collapse;width:100%;margin-bottom:25px}th,td{border:1px solid #d7d7d7;padding:8px;text-align:left;vertical-align:top}th{background:#eee}.summary{background:#fbfbfb}.summary th{background:#e9eef5}.row-number{width:60px;color:#777;font-weight:bold;background:#fafafa}.source{font-size:11px;color:#888;margin-top:4px}.sdt{background:#fff7d6}.ok{color:#187a31;font-weight:bold}.error{color:#b42318;font-weight:bold}.warning{color:#b26a00;font-weight:bold}.paragraph{padding:6px 0;border-bottom:1px solid #eee}.empty{color:#999;font-style:italic}.raw-section{margin-top:35px;border-top:3px solid #eee;padding-top:15px}.legend{padding:10px;margin-bottom:20px;background:#fff7d6;border:1px solid #eadb9c;border-radius:5px}</style></head><body><h1>Проверка Word-парсера</h1><div class="legend">Жёлтым цветом показаны значения, найденные внутри Word Content Control (w:sdt).</div>'''
    html += build_global_summary(documents)
    for document in documents:
        html += f'<div class="document"><div class="filename">{escape(document["filename"])}</div>{build_summary_table(document)}{build_employee_match_html(document)}<div class="raw-section"><h2>Исходная структура Word</h2><h3>Абзацы</h3>'
        if document["paragraphs"]:
            for paragraph in document["paragraphs"]:
                html += f'<div class="paragraph"><strong>P{paragraph["index"]}</strong>: {escape(paragraph["text"]).replace(chr(10), "<br>")}</div>'
        else:
            html += '<p class="empty">Текстовых абзацев не найдено</p>'
        html += "<h3>Таблицы</h3>"
        for table in document["tables"]:
            html += f'<h4>Таблица {table["table_index"]}</h4><table>'
            for row in table["rows"]:
                html += f'<tr><td class="row-number">R{row["row_index"]}</td>'
                for cell_index, cell in enumerate(row["cells"]):
                    css_class = ' class="sdt"' if cell["source_type"] == "sdt" else ""
                    text = escape(cell["text"]).replace(chr(10), "<br>") or '<span class="empty">пусто</span>'
                    html += f'<td{css_class}>{text}<div class="source">C{cell_index} · {cell["source_type"]}</div></td>'
                html += "</tr>"
            html += "</table>"
        html += "</div></div>"
    HTML_REPORT.write_text(html + "</body></html>", encoding="utf-8")


def main():
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    files = sorted(INPUT_FOLDER.glob("*.docx"))
    print(f"Найдено Word-файлов: {len(files)}")
    if not files:
        print("В папке input нет файлов .docx")
        return
    documents = []
    for file_path in files:
        print(f"Читаем: {file_path.name}")
        try:
            document_data = read_document(file_path)
            documents.append(document_data)
            inspector = document_data["business_data"]["inspector"]
            if inspector["status"] == "found":
                employee_match = document_data["business_data"]["employee_match"]
                position = (
                    employee_match["employee"]["position_name"]
                    if employee_match["status"] == "found"
                    else "не найден в БД"
                )
                print(
                    "  Специалист: "
                    f'{inspector["full_name"]["value"]} | {position}'
                )
            elif inspector["status"] == "ambiguous":
                print("  ВНИМАНИЕ: найдено несколько кандидатов на специалиста")
            else:
                print("  ОШИБКА: специалист не найден")
        except Exception as error:
            print(f"  ОШИБКА: {error}")
    build_html_report(documents)
    intermediate_report = build_intermediate_report(documents)

    print()
    print("Готово.")
    print(f"Диагностический отчёт: {HTML_REPORT.resolve()}")
    print(f"Промежуточный отчёт: {intermediate_report.resolve()}")


if __name__ == "__main__":
    main()
